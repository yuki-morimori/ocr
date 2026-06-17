# -*- coding: utf-8 -*-
"""給与素データ（設定駆動）。

同じ日報から、運賃請求（invoice）とは別に『ドライバー給与の素データ』を出す。
乗務員ごとに月集計し、手当（運行数×日当・走行km×単価 等）＋立替精算を計算する。

設定は各業種JSONの `payroll`（運送のみ同梱。他業種も同形式で追加可）:
  party_field … 集計の単位（乗務員）
  items …… 手当の定義 [{label, qty_field, rate, ...}]（qty_field="_count" は件数）
  reimburse … 立替精算（実費・給与とは別建て）

公開API:
    build(cfg, results, month=None, rates=None) -> list[paydata]
    render_html(cfg, paydata) -> str
    build_xlsx(cfg, paydatas, path) -> None   乗務員ごとに1シート
"""
from __future__ import annotations

from html import escape
from pathlib import Path
from typing import IO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_HEAD = PatternFill("solid", fgColor="1F2A3A")
_SUM = PatternFill("solid", fgColor="EAF2FF")


def _num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _yen(v) -> int:
    return int(round(v or 0))


def _in_month(date_val, month):
    return not month or (isinstance(date_val, str) and date_val.startswith(month))


def _qty(cfg: dict, slips: list[dict], field: str) -> float:
    if field == "_count":
        return len(slips)
    li = cfg["line_items"]["key"]
    total = 0.0
    for s in slips:
        # ヘッダ優先、無ければ明細合計
        v = _num(s.get(field))
        if v is not None:
            total += v
        else:
            total += sum(_num(r.get(field)) or 0 for r in s.get(li) or [])
    return total


def build(cfg: dict, results: list[dict], month: str | None = None,
          rates: dict | None = None) -> list[dict]:
    """確定データ → 乗務員ごとの給与素データ。月で絞り込み可。"""
    pc = cfg["payroll"]
    cur = pc.get("currency", "円")
    li = cfg["line_items"]["key"]
    rates = rates or {}

    groups: dict[str, list[dict]] = {}
    for s in results:
        if not _in_month(s.get("date"), month):
            continue
        party = s.get(pc["party_field"]) or "（氏名不明）"
        groups.setdefault(party, []).append(s)

    out = []
    for party in sorted(groups):
        slips = groups[party]
        lines = []
        for it in pc["items"]:
            qty = _qty(cfg, slips, it["qty_field"])
            rate = float(rates.get(it["label"], it.get("rate", 0)))
            lines.append({
                "label": it["label"], "qty": qty, "qty_label": it.get("qty_label", ""),
                "unit": it.get("unit", ""), "rate": rate, "amount": _yen(qty * rate),
            })
        allowance = sum(l["amount"] for l in lines)
        reimburse = 0
        reimburse_label = ""
        if pc.get("reimburse"):
            rb = pc["reimburse"]
            reimburse_label = rb["label"]
            reimburse = _yen(sum(
                sum(_num(r.get(rb["field"])) or 0 for r in s.get(li) or []) for s in slips
            ))
        out.append({
            "party": party, "party_label": pc.get("party_label", "対象"),
            "month": month or "全期間", "slip_count": len(slips),
            "lines": lines, "allowance_total": allowance,
            "reimburse_label": reimburse_label, "reimburse": reimburse,
            "total": allowance + reimburse, "currency": cur,
            "note": pc.get("note", ""),
        })
    return out


def _q(v):
    if v is None:
        return ""
    n = float(v)
    return f"{int(n):,}" if n == int(n) else f"{n:,.1f}"


def render_html(cfg: dict, pay: dict) -> str:
    cur = pay["currency"]
    rows = ""
    for l in pay["lines"]:
        rows += (f'<tr><td>{escape(l["label"])}</td>'
                 f'<td class="num">{_q(l["qty"])} {l["unit"]}</td>'
                 f'<td class="num">{_q(l["rate"])}{cur}</td>'
                 f'<td class="num">{_q(l["amount"])}{cur}</td></tr>')
    reimburse = ""
    if pay["reimburse_label"]:
        reimburse = (f'<tr><td colspan="3" class="total-label">{escape(pay["reimburse_label"])}</td>'
                     f'<td class="num">{_q(pay["reimburse"])}{cur}</td></tr>')
    note = f'<p class="form-note">※ {escape(pay["note"])}</p>' if pay["note"] else ""
    return f"""
<div class="paper invoice">
  <div class="paper-head"><div class="badge">{escape(cfg['display_name'])}</div>
    <h2>給与素データ</h2>
    <p class="subtitle">{escape(pay['party_label'])}: <b>{escape(pay['party'])}</b> ／
       対象: {escape(pay['month'])} ／ 運行 {pay['slip_count']} 件</p></div>
  <table class="lines">
    <thead><tr><th>項目</th><th>数量</th><th>単価</th><th>金額</th></tr></thead>
    <tbody>{rows}</tbody>
    <tfoot>
      <tr class="sumrow"><td colspan="3" class="total-label">手当計</td><td class="num total">{_q(pay['allowance_total'])}{cur}</td></tr>
      {reimburse}
      <tr class="sumrow grand"><td colspan="3" class="total-label">支給合計（手当＋立替精算）</td><td class="num total">{_q(pay['total'])}{cur}</td></tr>
    </tfoot>
  </table>
  {note}
</div>
""".strip()


def build_xlsx(cfg: dict, pays: list[dict], path: str | Path | IO) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    for i, pay in enumerate(pays):
        cur = pay["currency"]
        ws = wb.create_sheet((pay["party"][:25] or f"給与{i+1}").replace("/", "／")[:28])
        ws["A1"] = f"給与素データ（{cfg['display_name']}）"
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"] = f"{pay['party_label']}: {pay['party']} ／ 対象: {pay['month']} ／ 運行 {pay['slip_count']} 件"
        hr = 4
        for ci, h in enumerate(["項目", "数量", "単価", "金額"], 1):
            c = ws.cell(hr, ci, h)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = _HEAD
            c.alignment = Alignment(horizontal="center")
        r = hr + 1
        for l in pay["lines"]:
            ws.cell(r, 1, l["label"])
            ws.cell(r, 2, l["qty"])
            ws.cell(r, 3, l["rate"])
            ws.cell(r, 4, l["amount"])
            r += 1
        rows = [("手当計", pay["allowance_total"])]
        if pay["reimburse_label"]:
            rows.append((pay["reimburse_label"], pay["reimburse"]))
        rows.append(("支給合計（手当＋立替精算）", pay["total"]))
        for label, val in rows:
            ws.cell(r, 3, label).font = Font(bold=True)
            cell = ws.cell(r, 4, val)
            cell.font = Font(bold=True)
            cell.fill = _SUM
            r += 1
        for ci, w in enumerate([30, 14, 12, 14], 1):
            ws.column_dimensions[get_column_letter(ci)].width = w
    if not wb.sheetnames:
        wb.create_sheet("給与なし")
    wb.save(path)
