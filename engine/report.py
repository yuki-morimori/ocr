# -*- coding: utf-8 -*-
"""集計レポート（全業種共通・設定駆動）。

確定データ（複数伝票）を『軸（会社・車番・乗務員 等）』ごとに集計する内部コスト把握用。
請求書（顧客へ）とは別軸で、「どの会社が・どの車が・どれだけ経費/距離/重量か」を出す。

各業種JSONの `analytics` が軸(dimensions)と集計値(measures)を決める：

  運送 … 取引先別 / 車番別 / 乗務員別 に 立替経費・走行距離・運行数。会社×車番のクロスも。
  産廃 … 排出事業者別 / 車番別 に 正味重量・搬入回数。

公開API:
    aggregate(cfg, results, month=None) -> report
    render_html(cfg, report) -> str
    build_xlsx(cfg, report, path) -> None   軸ごとに1シート＋クロス
"""
from __future__ import annotations

from html import escape
from pathlib import Path
from typing import IO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_HEAD = PatternFill("solid", fgColor="1F2A3A")
_TOT = PatternFill("solid", fgColor="EAF2FF")
_UNKNOWN = "（不明）"


def _num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _in_month(date_val, month: str | None) -> bool:
    if not month:
        return True
    return isinstance(date_val, str) and date_val.startswith(month)


def _dim_value(slip: dict, key: str) -> str:
    v = slip.get(key)
    return v if v not in (None, "") else _UNKNOWN


def _slip_measure(cfg: dict, slip: dict, m: dict) -> float:
    if m.get("agg") == "count":
        return 1
    if m.get("source") == "line":
        li = cfg["line_items"]["key"]
        return sum(_num(r.get(m["key"])) or 0 for r in slip.get(li) or [])
    return _num(slip.get(m["key"])) or 0


def _sort_key(measures: list[dict]) -> str:
    """並べ替えの基準（最初の合計系メジャー。無ければ先頭）。"""
    for m in measures:
        if m.get("agg") != "count":
            return m["key"]
    return measures[0]["key"]


def aggregate(cfg: dict, results: list[dict], month: str | None = None) -> dict:
    """確定データ → 軸ごとの集計＋会社×車番クロス。月で絞り込み可。"""
    a = cfg["analytics"]
    measures = a["measures"]
    dims = a["dimensions"]
    rows_in = [s for s in results if _in_month(s.get("date"), month)]
    sort_k = _sort_key(measures)

    def _group(keyfunc):
        groups: dict = {}
        for s in rows_in:
            k = keyfunc(s)
            g = groups.setdefault(k, {m["key"]: 0.0 for m in measures})
            for m in measures:
                g[m["key"]] += _slip_measure(cfg, s, m)
        return groups

    out_dims = []
    for d in dims:
        groups = _group(lambda s, k=d["key"]: _dim_value(s, k))
        rows = [{"value": k, "measures": v} for k, v in groups.items()]
        rows.sort(key=lambda r: -r["measures"].get(sort_k, 0))
        totals = {m["key"]: sum(r["measures"][m["key"]] for r in rows) for m in measures}
        out_dims.append({"key": d["key"], "label": d["label"], "rows": rows, "totals": totals})

    cross = None
    if len(dims) >= 2:
        a0, b0 = dims[0], dims[1]
        groups = _group(lambda s: (_dim_value(s, a0["key"]), _dim_value(s, b0["key"])))
        rows = [{"a": k[0], "b": k[1], "measures": v} for k, v in groups.items()]
        rows.sort(key=lambda r: (r["a"], -r["measures"].get(sort_k, 0)))
        cross = {"a_label": a0["label"], "b_label": b0["label"], "rows": rows}

    return {
        "month": month or "全期間",
        "slip_count": len(rows_in),
        "measures": [{"key": m["key"], "label": m["label"], "unit": m.get("unit", "")} for m in measures],
        "dimensions": out_dims,
        "cross": cross,
    }


# ── 表示（HTML）────────────────────────────────────────────────────────────────
def _fmt(v, unit: str) -> str:
    try:
        n = float(v)
    except (TypeError, ValueError):
        return escape(str(v))
    s = f"{int(n):,}" if n == int(n) else f"{n:,.1f}"
    return f"{s}{unit}"


def _table(measures, head_labels, body_rows) -> str:
    th = "".join(f"<th>{escape(h)}</th>" for h in head_labels)
    th += "".join(f'<th>{escape(m["label"])}</th>' for m in measures)
    return (f'<table class="lines"><thead><tr>{th}</tr></thead>'
            f'<tbody>{body_rows}</tbody></table>')


def render_html(cfg: dict, rep: dict) -> str:
    measures = rep["measures"]
    blocks = [f'<p class="subtitle">対象期間: {escape(rep["month"])} ／ 伝票 {rep["slip_count"]} 枚</p>']
    for d in rep["dimensions"]:
        body = ""
        for r in d["rows"]:
            cells = f'<td>{escape(str(r["value"]))}</td>'
            cells += "".join(f'<td class="num">{_fmt(r["measures"][m["key"]], m["unit"])}</td>' for m in measures)
            body += f"<tr>{cells}</tr>"
        tot = f'<td class="total-label">合計</td>' + "".join(
            f'<td class="num total">{_fmt(d["totals"][m["key"]], m["unit"])}</td>' for m in measures)
        body += f'<tr class="sumrow">{tot}</tr>'
        blocks.append(f'<h3 class="rep-h">{escape(d["label"])}別</h3>' + _table(measures, [d["label"]], body))
    if rep["cross"]:
        cx = rep["cross"]
        body = ""
        for r in cx["rows"]:
            cells = f'<td>{escape(str(r["a"]))}</td><td>{escape(str(r["b"]))}</td>'
            cells += "".join(f'<td class="num">{_fmt(r["measures"][m["key"]], m["unit"])}</td>' for m in measures)
            body += f"<tr>{cells}</tr>"
        blocks.append(f'<h3 class="rep-h">{escape(cx["a_label"])} × {escape(cx["b_label"])}（クロス集計）</h3>'
                      + _table(measures, [cx["a_label"], cx["b_label"]], body))
    inner = "\n".join(blocks)
    return f"""
<div class="paper">
  <div class="paper-head"><div class="badge">{escape(cfg['display_name'])}</div>
    <h2>集計レポート</h2></div>
  {inner}
</div>
""".strip()


# ── 出力（Excel：軸ごとに1シート＋クロス）─────────────────────────────────────
def _sheet(wb, title: str):
    ws = wb.create_sheet(title[:28] or "集計")
    return ws


def _write_table(ws, head_labels, measures, rows, totals=None):
    for ci, h in enumerate(head_labels + [m["label"] for m in measures], 1):
        c = ws.cell(1, ci, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = _HEAD
        c.alignment = Alignment(horizontal="center")
    r = 2
    for row in rows:
        for ci, v in enumerate(row, 1):
            ws.cell(r, ci, v)
        r += 1
    if totals is not None:
        ws.cell(r, len(head_labels), "合計").font = Font(bold=True)
        for j, m in enumerate(measures):
            cell = ws.cell(r, len(head_labels) + 1 + j, totals[m["key"]])
            cell.font = Font(bold=True)
            cell.fill = _TOT
    for ci in range(1, len(head_labels) + len(measures) + 1):
        ws.column_dimensions[get_column_letter(ci)].width = 16


def build_xlsx(cfg: dict, rep: dict, path: str | Path | IO) -> None:
    measures = rep["measures"]
    wb = Workbook()
    wb.remove(wb.active)
    for d in rep["dimensions"]:
        ws = _sheet(wb, f"{d['label']}別")
        rows = [[r["value"]] + [r["measures"][m["key"]] for m in measures] for r in d["rows"]]
        _write_table(ws, [d["label"]], measures, rows, d["totals"])
    if rep["cross"]:
        cx = rep["cross"]
        ws = _sheet(wb, f"{cx['a_label']}×{cx['b_label']}")
        rows = [[r["a"], r["b"]] + [r["measures"][m["key"]] for m in measures] for r in cx["rows"]]
        _write_table(ws, [cx["a_label"], cx["b_label"]], measures, rows)
    if not wb.sheetnames:
        wb.create_sheet("集計なし")
    wb.save(path)
