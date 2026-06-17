# -*- coding: utf-8 -*-
"""確認シート（全業種共通）。

エスカレーションの終点＝『人間の目視』を担う、最終確認・訂正の砦。

    build(cfg, results, path)  構造化結果（複数伝票）→ Excel確認シート
    read(cfg, path)            人が訂正したExcel → 構造化結果（帳票/請求書へ流す）

要人間確認(needs_human_review)＝赤、低自信度(low)＝黄でセルを色分けし、
人はそこだけ原本を見て訂正する。訂正後に read で読み戻すとループが閉じる：

    画像 → OCR(エスカレーション) → 確認シート → 人が訂正 → read → 帳票/請求書

列の順序は build / read で同じ _columns(cfg) から作るので、位置でキーに対応づく。
"""
from __future__ import annotations

import json
from pathlib import Path
from typing import IO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import core, learning

_META_SHEET = "_ocr_meta"  # 読み戻し時に元のOCR値と突き合わせて学習するための隠しシート

# セルのハイライト色
_RED = PatternFill("solid", fgColor="F6C6C0")      # 要人間確認
_YELLOW = PatternFill("solid", fgColor="FCE2A8")   # 低自信度
_HEAD = PatternFill("solid", fgColor="1F2A3A")     # 見出し行

_HEADER_ROW = 4  # 列見出しの行（1=タイトル, 2=凡例, 3=空け, 4=見出し, 5〜=データ）


def _columns(cfg: dict) -> list[tuple[str, str, str]]:
    """(キー, 見出しラベル, 種別) の列定義。build/read で共有し位置対応させる。"""
    cols: list[tuple[str, str, str]] = [("__slip", "伝票#", "meta")]
    for f in cfg["header_fields"]:
        cols.append((f["key"], f["label"], "header"))
    for c in cfg["line_items"]["columns"]:
        cols.append((c["key"], c["label"], "line"))
    cols += [("__confirm", "確認", "meta"), ("__model", "確定モデル", "meta"), ("__memo", "メモ", "meta")]
    return cols


def build(cfg: dict, results: list[dict], path: str | Path | IO) -> None:
    """構造化結果（複数伝票）から Excel 確認シートを生成する。"""
    cols = _columns(cfg)
    wb = Workbook()
    ws = wb.active
    ws.title = "確認シート"

    ws["A1"] = f"{cfg['display_name']} 確認シート"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = ("赤=要人間確認 / 黄=低自信度。色のついたセルだけ原本を見て訂正し、"
                "『確認』列をOKにしてください。数字は推測で埋めない。")
    ws["A2"].font = Font(color="8A5A00")

    for ci, (_key, label, _kind) in enumerate(cols, 1):
        cell = ws.cell(_HEADER_ROW, ci, label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEAD
        cell.alignment = Alignment(horizontal="center")

    r = _HEADER_ROW + 1
    for si, res in enumerate(results, 1):
        review = set(res.get("needs_human_review") or [])
        low = {k for k, v in (res.get("confidence") or {}).items() if str(v).lower() == "low"}
        model = (res.get("_escalation") or {}).get("final_model") or ""
        flagged = bool(review or low)
        rows = res.get(cfg["line_items"]["key"]) or [{}]
        for line in rows:
            for ci, (key, _label, kind) in enumerate(cols, 1):
                if key == "__slip":
                    val = si
                elif key == "__confirm":
                    val = "要確認" if flagged else "OK"
                elif key == "__model":
                    val = model
                elif key == "__memo":
                    val = None
                elif kind == "header":
                    val = res.get(key)
                else:
                    val = line.get(key)
                cell = ws.cell(r, ci, val)
                if key in review:
                    cell.fill = _RED
                elif key in low:
                    cell.fill = _YELLOW
            r += 1

    ws.freeze_panes = f"A{_HEADER_ROW + 1}"
    for ci, (_key, label, _kind) in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(ci)].width = max(11, len(label) + 3)

    # 元のOCR値を隠しシートに埋め込む（読み戻し時に人手訂正との差分を学習するため）
    meta = wb.create_sheet(_META_SHEET)
    meta.sheet_state = "hidden"
    meta["A1"] = "自動学習用データです。編集・削除しないでください。"
    meta["A2"] = json.dumps({"industry": cfg["id"], "original": results}, ensure_ascii=False)

    wb.save(path)


def read(cfg: dict, path: str | Path | IO, learn: bool = True) -> list[dict]:
    """人が訂正した確認シートを読み戻し、伝票ごとの構造化結果に戻す。

    『伝票#』でグループ化し、ヘッダ＝そのグループ先頭行、明細＝各行から再構成する。
    訂正済みなので confidence / unreadable は空（フラグは消える）。

    learn=True なら、隠しシートの元OCR値と突き合わせて学習を更新する
    （＝読み戻すたびに次回以降が賢くなる）。
    """
    cols = _columns(cfg)
    keys = [c[0] for c in cols]
    header_keys = [k for k, _l, kind in cols if kind == "header"]
    line_keys = [k for k, _l, kind in cols if kind == "line"]
    li_key = cfg["line_items"]["key"]

    wb = load_workbook(path, data_only=True)
    ws = wb["確認シート"] if "確認シート" in wb.sheetnames else wb.worksheets[0]
    slips: dict = {}
    order: list = []
    for row in ws.iter_rows(min_row=_HEADER_ROW + 1, values_only=True):
        if row is None or all(v is None for v in row):
            continue
        rec = dict(zip(keys, row))
        slip = rec.get("__slip")
        if slip is None:
            continue
        if slip not in slips:
            slips[slip] = {k: rec.get(k) for k in header_keys}
            slips[slip][li_key] = []
            order.append(slip)
        line = {k: rec.get(k) for k in line_keys}
        if any(v not in (None, "") for v in line.values()):
            slips[slip][li_key].append(line)

    corrected = [core._normalize(cfg, slips[s]) for s in order]

    # 学習：隠しシートの元OCR値と突き合わせて訂正・語彙を蓄積
    if learn and _META_SHEET in wb.sheetnames:
        try:
            meta = json.loads(wb[_META_SHEET]["A2"].value or "{}")
            original = meta.get("original") or []
            if original:
                learning.learn_from(cfg, original, corrected)
        except Exception:  # noqa: BLE001 — 学習失敗で読み戻しは止めない
            pass

    return corrected
