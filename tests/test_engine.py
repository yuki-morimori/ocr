# -*- coding: utf-8 -*-
"""コア＋設定＋帳票の最小テスト（APIキー不要＝モード非依存で動く）。

実行: cd ocr && python -m pytest -q
"""
import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from engine import core, forms, invoice, learning, registry, report, review_sheet  # noqa: E402

from openpyxl import load_workbook  # noqa: E402


def test_industries_registered():
    ids = {it["id"] for it in registry.list_industries()}
    assert {"sanpai", "unso"} <= ids


def test_each_config_builds_prompt_with_4_elements():
    """各業種設定から、文脈宣言・用語辞書・出力の型・補完抑制を含むプロンプトが組める。"""
    for it in registry.list_industries():
        cfg = registry.get(it["id"])
        sp = core.build_system_prompt(cfg)
        assert cfg["context_declaration"] in sp        # 文脈宣言
        assert "用語辞書" in sp                          # 用語辞書
        assert "出力するJSONの型" in sp                  # 出力の型
        assert "補完抑制ルール" in sp                    # 補完抑制ルール
        # トップレベルのキー順が固定で指示されている
        for f in cfg["header_fields"]:
            assert f'"{f["key"]}"' in sp


def test_normalize_fills_and_orders_keys():
    cfg = registry.get("sanpai")
    raw = {"emitter": "テスト商事", "items": [{"item": "混合", "net": 1000}], "extra": "捨てる"}
    out = core._normalize(cfg, raw)
    # ヘッダの全キーが存在し、未記入は None
    for f in cfg["header_fields"]:
        assert f["key"] in out
    assert out["emitter"] == "テスト商事"
    assert out["date"] is None
    # 明細は列キーに正規化される
    assert out["items"][0]["item"] == "混合"
    assert out["items"][0]["net"] == 1000
    assert "tare" in out["items"][0]
    assert "extra" not in out


def test_extract_mock_runs_without_apikey(monkeypatch):
    """APIキーが無くてもモックで構造化結果が返る（デモが止まらない）。"""
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    cfg = registry.get("unso")
    data = core.extract("dummy.jpg", cfg)
    assert data["driver"] is not None
    assert isinstance(data["trips"], list) and data["trips"]


def test_model_ladder_and_capabilities():
    """段位はHaiku→Sonnet→Opus。Haikuはeffort非対応なので付けない。"""
    assert core.resolve_ladder() == ["haiku", "sonnet", "opus"]
    assert core.MODELS["haiku"]["id"] == "claude-haiku-4-5"
    assert core.MODELS["sonnet"]["id"] == "claude-sonnet-4-6"
    assert core.MODELS["opus"]["id"] == "claude-opus-4-8"
    assert core.MODELS["haiku"]["effort"] is None          # Haiku 4.5 は effort 非対応
    assert core.MODELS["haiku"]["thinking"] is False        # 最安・抽出に十分
    assert core.MODELS["opus"]["thinking"] is True


def test_ladder_override(monkeypatch):
    monkeypatch.setenv("OCR_LADDER", "haiku,sonnet")
    assert core.resolve_ladder() == ["haiku", "sonnet"]


def test_problem_fields_low_confidence_and_critical_empty():
    cfg = registry.get("sanpai")
    # low自信度 → 問題、medium → 問題ではない
    d1 = core._normalize(cfg, {"confidence": {"emitter": "low", "site": "medium"}})
    p1 = core.problem_fields(cfg, d1)
    assert "emitter" in p1 and "site" not in p1
    # 重要項目(net)が空の明細 → 申告が無くても問題（保守側）
    d2 = core._normalize(cfg, {"items": [{"item": "混合", "net": None}]})
    assert "net" in core.problem_fields(cfg, d2)
    # 重要項目が埋まっていれば問題なし
    d3 = core._normalize(cfg, {"emitter": "A社", "items": [{"item": "混合", "gross": 100, "tare": 30, "net": 70}]})
    assert core.problem_fields(cfg, d3) == set()


def test_mock_escalation_trace_present():
    # unso は low_first_pass を持つので Haiku→Sonnet の擬似トレースが出る
    cfg = registry.get("unso")
    data = core.extract("dummy.jpg", cfg)
    esc = data["_escalation"]
    assert esc["mock"] is True
    assert [s["tier"] for s in esc["steps"]] == ["haiku", "sonnet"]
    assert esc["escalated"] is True
    assert data["needs_human_review"] == []         # Sonnetで解決
    # sanpai は一発で確定（Haikuのみ）
    s = core.extract("dummy.jpg", registry.get("sanpai"))["_escalation"]
    assert [x["tier"] for x in s["steps"]] == ["haiku"]


def test_cost_estimate_matches_handoff():
    # 引き継ぎメモの実測: Haiku≈0.75円, Sonnet≈2.25円, Opus≈3.75円/枚（1ドル150円）
    assert abs(core._est_cost("haiku") - 0.75) < 0.01
    assert abs(core._est_cost("sonnet") - 2.25) < 0.01
    assert abs(core._est_cost("opus") - 3.75) < 0.01


def test_prompt_declares_confidence_and_critical():
    cfg = registry.get("unso")
    sp = core.build_system_prompt(cfg)
    assert "自信度の申告" in sp
    assert "confidence" in sp
    assert "★" in sp  # 重要項目(critical)の印
    assert '"confidence"' in sp and '"needs_human_review"' not in sp  # confidenceは出力キー


def test_render_form_sums_numeric_columns():
    cfg = registry.get("sanpai")
    data = core.extract("dummy.jpg", cfg)
    html = forms.render_form(cfg, data)
    assert cfg["form"]["title"] in html
    assert "合計" in html  # sum:true 列の合計行
    assert "正味重量" in html


def test_unknown_industry_raises():
    try:
        registry.get("no-such")
        assert False, "should raise"
    except KeyError:
        pass


def _hl(cfg, d):
    """比較用にヘッダ＋明細だけ取り出す。"""
    li = cfg["line_items"]["key"]
    return ({f["key"]: d.get(f["key"]) for f in cfg["header_fields"]}, d.get(li))


def test_review_sheet_roundtrip_preserves_values():
    """確認シートを出力→読み戻して、ヘッダ・明細が一致する。"""
    cfg = registry.get("sanpai")
    results = [core.extract("a.jpg", cfg), core.extract("b.jpg", cfg)]
    bio = io.BytesIO()
    review_sheet.build(cfg, results, bio)
    bio.seek(0)
    back = review_sheet.read(cfg, bio, learn=False)
    assert len(back) == len(results)
    for src, got in zip(results, back):
        assert _hl(cfg, src) == _hl(cfg, got)


def test_review_sheet_flags_human_review(tmp_path):
    """要人間確認の伝票は『要確認』、それ以外は『OK』が入る。"""
    cfg = registry.get("unso")
    res = core.extract("x.jpg", cfg)
    res["needs_human_review"] = ["expense_amount"]  # 強制的に要確認に
    p = tmp_path / "review.xlsx"
    review_sheet.build(cfg, [res], p)
    wb = load_workbook(p)
    ws = wb["確認シート"]
    confirm_col = [c[1] for c in review_sheet._columns(cfg)].index("確認") + 1
    vals = [ws.cell(r, confirm_col).value for r in range(5, ws.max_row + 1)]
    assert "要確認" in vals


def test_learning_makes_it_smarter(monkeypatch, tmp_path):
    """訂正を読み戻すと、誤読の訂正と語彙が蓄積され、次回プロンプトに載る。"""
    monkeypatch.setenv("OCR_LEARN_DIR", str(tmp_path))
    cfg = registry.get("sanpai")
    # OCRが「丸北健設」と誤読 → 人が「丸北建設」に訂正
    original = [core._normalize(cfg, {"emitter": "丸北健設", "items": [{"item": "混合廃棄物"}]})]
    corrected = [core._normalize(cfg, {"emitter": "丸北建設", "items": [{"item": "混合廃棄物"}]})]
    stats = learning.learn_from(cfg, original, corrected)
    assert stats["new_corrections"] == 1
    assert stats["total_vocabulary"] >= 2  # 排出事業者＋品目

    block = learning.as_prompt_block(cfg)
    assert "丸北健設" in block and "丸北建設" in block      # 誤読→正の訂正
    assert "混合廃棄物" in block                            # 確認済み語彙

    # 次回の system プロンプトに学習が差し込まれている＝賢くなっている
    sp = core.build_system_prompt(cfg)
    assert "現場で確定した学習メモ" in sp
    assert "丸北建設" in sp

    # もう一度同じ訂正 → 回数が増える（累積する）
    learning.learn_from(cfg, original, corrected)
    again = learning.load(cfg)["corrections"][0]
    assert again["count"] == 2


def test_learning_loop_through_sheet(monkeypatch, tmp_path):
    """確認シート経由でも学習が回る（隠しシートの元値と差分を取る）。"""
    monkeypatch.setenv("OCR_LEARN_DIR", str(tmp_path))
    cfg = registry.get("unso")
    res = core.extract("r.jpg", cfg)  # mock。driver="濱崎"
    p = tmp_path / "review.xlsx"
    review_sheet.build(cfg, [res], p)
    # 人が乗務員名を訂正（濱崎→浜崎）
    wb = load_workbook(p)
    ws = wb["確認シート"]
    driver_col = [c[0] for c in review_sheet._columns(cfg)].index("driver") + 1
    ws.cell(5, driver_col).value = "浜崎"
    wb.save(p)
    before = learning.summary(cfg)["corrections"]
    review_sheet.read(cfg, p, learn=True)
    after = learning.summary(cfg)["corrections"]
    assert after == before + 1


def test_invoice_sanpai_weight_pricing():
    """産廃: 排出事業者ごとに 正味重量×区分単価（混合は高単価）で集計。"""
    cfg = registry.get("sanpai")
    slips = [core._normalize(cfg, {
        "date": "2026-06-10", "emitter": "A社",
        "items": [{"item": "混合廃棄物", "category": "混合", "net": 1000},
                  {"item": "がれき類", "category": "単品", "net": 2000}],
    })]
    invs = invoice.build_invoices(cfg, slips, month="2026-06")
    assert len(invs) == 1
    inv = invs[0]
    assert inv["party"] == "A社"
    amounts = {l["desc"]: l["amount"] for l in inv["lines"]}
    assert amounts["混合廃棄物（混合）"] == 25000   # 1000kg × 25
    assert amounts["がれき類（単品）"] == 24000     # 2000kg × 12
    assert inv["subtotal"] == 49000
    assert inv["tax"] == 4900 and inv["total"] == 53900


def test_invoice_groups_by_party_and_month():
    cfg = registry.get("sanpai")
    slips = [
        core._normalize(cfg, {"date": "2026-06-01", "emitter": "A社", "items": [{"category": "混合", "net": 100}]}),
        core._normalize(cfg, {"date": "2026-06-02", "emitter": "B社", "items": [{"category": "混合", "net": 200}]}),
        core._normalize(cfg, {"date": "2026-05-30", "emitter": "A社", "items": [{"category": "混合", "net": 999}]}),
    ]
    invs = invoice.build_invoices(cfg, slips, month="2026-06")
    parties = {i["party"]: i for i in invs}
    assert set(parties) == {"A社", "B社"}                # 5月分は除外
    assert parties["A社"]["subtotal"] == 2500            # 100kg×25 のみ（5月の999は対象外）


def test_invoice_unso_trip_and_charter():
    """運送: 走行×運賃＋立替。傭車便は別仕分け。"""
    cfg = registry.get("unso")
    slips = [
        core._normalize(cfg, {"date": "2026-06-05", "client": "X物流", "total_distance": 100,
                              "trips": [{"charter": "自社", "expense_amount": 500}]}),
        core._normalize(cfg, {"date": "2026-06-06", "client": "X物流", "total_distance": 50,
                              "trips": [{"charter": "傭車", "expense_amount": 800}]}),
    ]
    inv = invoice.build_invoices(cfg, slips, month="2026-06")[0]
    fare = next(l for l in inv["lines"] if "運賃" in l["desc"])
    assert fare["amount"] == 12000                       # 100km × 120（自社のみ）
    assert any("立替" in l["desc"] for l in inv["lines"])
    assert len(inv["charter_lines"]) == 1                # 傭車は別仕分け
    assert inv["charter_lines"][0]["amount"] == 800


def test_invoice_xlsx_writes_sheet_per_party(tmp_path):
    cfg = registry.get("sanpai")
    slips = [core.extract("a.jpg", cfg)]
    invs = invoice.build_invoices(cfg, slips)
    p = tmp_path / "invoice.xlsx"
    invoice.build_xlsx(cfg, invs, p)
    from openpyxl import load_workbook
    wb = load_workbook(p)
    assert len(wb.sheetnames) == len(invs) >= 1


def test_both_industries_have_billing():
    for it in registry.list_industries():
        assert "billing" in registry.get(it["id"])


def test_report_aggregates_by_company_and_vehicle():
    """会社別・車番別・クロスで 立替経費/走行距離/件数 を集計する。"""
    cfg = registry.get("unso")
    N = core._normalize
    slips = [
        N(cfg, {"date": "2026-06-01", "client": "A物流", "vehicle_no": "車1", "driver": "甲",
                "total_distance": 100, "trips": [{"expense_amount": 1000}, {"expense_amount": 500}]}),
        N(cfg, {"date": "2026-06-02", "client": "A物流", "vehicle_no": "車1", "driver": "甲",
                "total_distance": 50, "trips": [{"expense_amount": 200}]}),
        N(cfg, {"date": "2026-06-03", "client": "B運輸", "vehicle_no": "車2", "driver": "乙",
                "total_distance": 300, "trips": [{"expense_amount": 800}]}),
    ]
    rep = report.aggregate(cfg, slips, month="2026-06")
    by = {d["label"]: d for d in rep["dimensions"]}
    assert {"取引先", "車番", "乗務員"} <= set(by)
    car1 = next(r for r in by["車番"]["rows"] if r["value"] == "車1")
    assert car1["measures"]["expense_amount"] == 1700   # 1000+500+200
    assert car1["measures"]["total_distance"] == 150
    assert car1["measures"]["_count"] == 2
    # クロス（取引先×車番）
    assert rep["cross"]["a_label"] == "取引先" and rep["cross"]["b_label"] == "車番"
    a_car1 = next(r for r in rep["cross"]["rows"] if r["a"] == "A物流" and r["b"] == "車1")
    assert a_car1["measures"]["expense_amount"] == 1700


def test_report_month_filter_and_unknown():
    cfg = registry.get("unso")
    N = core._normalize
    slips = [
        N(cfg, {"date": "2026-06-01", "client": "A", "vehicle_no": "車1", "trips": [{"expense_amount": 100}]}),
        N(cfg, {"date": "2026-05-01", "client": "A", "vehicle_no": "車1", "trips": [{"expense_amount": 999}]}),
        N(cfg, {"date": "2026-06-02", "vehicle_no": None, "trips": [{"expense_amount": 50}]}),
    ]
    rep = report.aggregate(cfg, slips, month="2026-06")
    assert rep["slip_count"] == 2                       # 5月分は除外
    veh = {r["value"]: r for r in next(d for d in rep["dimensions"] if d["label"] == "車番")["rows"]}
    assert "（不明）" in veh                             # 車番なしは（不明）に集計


def test_report_xlsx_has_sheet_per_dimension_plus_cross(tmp_path):
    cfg = registry.get("unso")
    slips = [core.extract("a.jpg", cfg)]
    rep = report.aggregate(cfg, slips)
    p = tmp_path / "report.xlsx"
    report.build_xlsx(cfg, rep, p)
    from openpyxl import load_workbook
    names = load_workbook(p).sheetnames
    # 3軸 + クロス1
    assert len(names) == len(rep["dimensions"]) + 1


def test_both_industries_have_analytics():
    for it in registry.list_industries():
        assert "analytics" in registry.get(it["id"])
